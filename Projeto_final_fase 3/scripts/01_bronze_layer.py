from pathlib import Path
from openpyxl import load_workbook
from pyspark.sql import functions as F
from pyspark.sql import types as T
from spark_utils import get_spark, write_single_csv

spark = get_spark("BronzeLayer")
raw_file = Path("data/raw/dados_credito.xlsx")
if not raw_file.exists():
    raise FileNotFoundError(f"Arquivo não encontrado: {raw_file}")

schema = T.StructType([
    T.StructField("CODIGO_CLIENTE", T.LongType(), True),
    T.StructField("UF", T.StringType(), True),
    T.StructField("IDADE", T.IntegerType(), True),
    T.StructField("ESCOLARIDADE", T.StringType(), True),
    T.StructField("ESTADO_CIVIL", T.StringType(), True),
    T.StructField("QT_FILHOS", T.IntegerType(), True),
    T.StructField("CASA_PROPRIA", T.StringType(), True),
    T.StructField("QT_IMOVEIS", T.IntegerType(), True),
    T.StructField("VL_IMOVEIS", T.DoubleType(), True),
    T.StructField("OUTRA_RENDA", T.StringType(), True),
    T.StructField("OUTRA_RENDA_VALOR", T.DoubleType(), True),
    T.StructField("TEMPO_ULTIMO_EMPREGO_MESES", T.IntegerType(), True),
    T.StructField("TRABALHANDO_ATUALMENTE", T.StringType(), True),
    T.StructField("ULTIMO_SALARIO", T.DoubleType(), True),
    T.StructField("QT_CARROS", T.IntegerType(), True),
    T.StructField("VALOR_TABELA_CARROS", T.DoubleType(), True),
    T.StructField("SCORE", T.DoubleType(), True)
])

for folder in [Path("data/bronze"), Path("data/silver"), Path("data/gold")]:
    folder.mkdir(parents=True, exist_ok=True)
print("Estrutura de diretórios verificada.")

wb = load_workbook(raw_file, data_only=True)
sheet = wb.active
rows_iter = sheet.iter_rows(values_only=True)
headers = [str(value).strip() for value in next(rows_iter)]
records = [dict(zip(headers, row)) for row in rows_iter if any(row)]

blank_markers = {"", "None", "NULL"}

def cast_value(value, data_type):
    if value is None:
        return None
    if isinstance(value, str):
        cleaned = value.strip()
        if cleaned in blank_markers:
            return None
        value = cleaned
    try:
        if isinstance(data_type, T.IntegralType):
            return int(float(value))
        if isinstance(data_type, T.FractionalType):
            return float(value)
    except (ValueError, TypeError):
        return None
    return str(value)

clean_records = []
for record in records:
    sanitized = {}
    for field in schema:
        sanitized[field.name] = cast_value(record.get(field.name), field.dataType)
    clean_records.append(sanitized)

df = spark.createDataFrame(clean_records, schema=schema)
print(f"Dimensão inicial do dataset: ({df.count()}, {len(df.columns)})")
df.printSchema()

df.describe().show(truncate=False)
null_overview = df.select([F.count(F.when(F.col(c).isNull(), 1)).alias(c) for c in df.columns])
null_overview.show(truncate=False)

df = (
    df.withColumn("DATA_UPLOAD", F.current_timestamp())
      .withColumn("ARQUIVO_FONTE", F.lit(raw_file.name))
)

bronze_path = "data/bronze/dados_brutos.csv"
write_single_csv(df, bronze_path)
print(f">>> Bronze salvo em: {bronze_path}")
df.show(5, truncate=False)
