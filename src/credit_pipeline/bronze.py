"""Camada Bronze: ingestão do arquivo bruto com schema explícito e metadados de linhagem."""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook
from pyspark.sql import DataFrame
from pyspark.sql import functions as F
from pyspark.sql import types as T

from credit_pipeline.config import Paths
from credit_pipeline.io import write_layer
from credit_pipeline.spark import get_spark

logger = logging.getLogger(__name__)

# Contrato de entrada: qualquer coluna ausente na fonte vira uma coluna nula do tipo esperado.
SCHEMA = T.StructType([
    T.StructField("CODIGO_CLIENTE", T.LongType()),
    T.StructField("UF", T.StringType()),
    T.StructField("IDADE", T.IntegerType()),
    T.StructField("ESCOLARIDADE", T.StringType()),
    T.StructField("ESTADO_CIVIL", T.StringType()),
    T.StructField("QT_FILHOS", T.IntegerType()),
    T.StructField("CASA_PROPRIA", T.StringType()),
    T.StructField("QT_IMOVEIS", T.IntegerType()),
    T.StructField("VL_IMOVEIS", T.DoubleType()),
    T.StructField("OUTRA_RENDA", T.StringType()),
    T.StructField("OUTRA_RENDA_VALOR", T.DoubleType()),
    T.StructField("TEMPO_ULTIMO_EMPREGO_MESES", T.IntegerType()),
    T.StructField("TRABALHANDO_ATUALMENTE", T.StringType()),
    T.StructField("ULTIMO_SALARIO", T.DoubleType()),
    T.StructField("QT_CARROS", T.IntegerType()),
    T.StructField("VALOR_TABELA_CARROS", T.DoubleType()),
    T.StructField("SCORE", T.DoubleType()),
])

BLANK_MARKERS = {"", "None", "NULL"}


def cast_value(value, data_type: T.DataType):
    """Aplica o tipo do schema, tratando marcadores de vazio e valores não numéricos como nulos."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value in BLANK_MARKERS:
            return None
    try:
        if isinstance(data_type, T.IntegralType):
            return int(float(value))
        if isinstance(data_type, T.FractionalType):
            return float(value)
    except (ValueError, TypeError):
        return None
    return str(value)


def read_excel_records(raw_file: Path) -> list[dict]:
    sheet = load_workbook(raw_file, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(rows)]
    return [dict(zip(headers, row, strict=False)) for row in rows if any(v is not None for v in row)]


def extract(raw_file: Path) -> DataFrame:
    """Lê o Excel de origem e devolve um DataFrame aderente ao SCHEMA, com colunas de linhagem."""
    if not raw_file.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {raw_file}")

    records = [
        tuple(cast_value(record.get(field.name), field.dataType) for field in SCHEMA)
        for record in read_excel_records(raw_file)
    ]
    return (
        get_spark()
        .createDataFrame(records, schema=SCHEMA)
        .withColumn("DATA_UPLOAD", F.current_timestamp())
        .withColumn("ARQUIVO_FONTE", F.lit(raw_file.name))
    )


def run(paths: Paths) -> DataFrame:
    df = extract(paths.raw_file)
    write_layer(df, paths.bronze_path)
    logger.info("Bronze: %d colunas -> %s", len(df.columns), paths.bronze_path)
    return df
